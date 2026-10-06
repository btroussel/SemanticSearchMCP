"""Repository discovery, ignore handling, and confined source reads."""
from __future__ import annotations

import hashlib
import os
import subprocess
from pathlib import Path

import pathspec

EXTENSIONS = {".py", ".pyi", ".js", ".jsx", ".ts", ".tsx", ".mjs", ".cjs",
              ".go", ".rs", ".java", ".c", ".h", ".cpp", ".hpp", ".cc",
              ".cs", ".rb", ".php", ".swift", ".kt", ".scala", ".sh",
              ".sql", ".md", ".rst", ".txt", ".toml", ".yaml", ".yml",
              ".json", ".html", ".css", ".vue", ".svelte"}
EXCLUDED = {".git", ".venv", "venv", "node_modules", "vendor",
            ".code-search", "__pycache__", ".pytest_cache", "dist", "build",
            ".next", ".mypy_cache", ".ruff_cache", "coverage", "target"}
MAX_BYTES = 512_000
MEDIA_MAX_BYTES = 20_000_000
DOCUMENT_EXTENSIONS = {".md", ".rst", ".txt", ".pdf", ".docx", ".csv", ".log"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}


def digest(text: str) -> str:
    return hashlib.sha256(text.encode()).hexdigest()


class Repository:
    def __init__(self, root: Path, kinds: list[str] | None = None):
        self.root = root.expanduser().resolve()
        if not self.root.is_dir():
            raise ValueError(f"Repository does not exist: {self.root}")
        self.kinds = set(kinds or ["code", "documents"])
        self.extensions = set(EXTENSIONS) if kinds is None else set()
        if kinds is not None:
            if "code" in self.kinds:
                self.extensions.update(EXTENSIONS - DOCUMENT_EXTENSIONS)
            if "documents" in self.kinds:
                self.extensions.update(DOCUMENT_EXTENSIONS)
            if "images" in self.kinds:
                self.extensions.update(IMAGE_EXTENSIONS)

    def kind(self, relative: str) -> str:
        suffix = Path(relative).suffix.lower()
        return "images" if suffix in IMAGE_EXTENSIONS else "documents" if suffix in DOCUMENT_EXTENSIONS else "code"

    def supported(self, relative: str) -> bool:
        p = Path(relative)
        return p.suffix.lower() in self.extensions or ("code" in self.kinds and p.name in {"Dockerfile", "Makefile"})

    def size_limit(self, relative: str) -> int:
        return MEDIA_MAX_BYTES if Path(relative).suffix.lower() in IMAGE_EXTENSIONS | {".pdf", ".docx"} else MAX_BYTES

    def _spec(self, directory: Path) -> pathspec.PathSpec:
        patterns = []
        for name in (".gitignore", ".code-searchignore", ".cursorignore"):
            file = directory / name
            if file.is_file() and not file.is_symlink():
                patterns.extend(file.read_text(errors="replace").splitlines())
        return pathspec.PathSpec.from_lines("gitignore", patterns)

    def ignored(self, relative: str) -> bool:
        path = Path(relative)
        if path.is_absolute() or ".." in path.parts:
            return True
        if any(p in EXCLUDED for p in path.parts):
            return True
        if path.name.startswith(".env") or path.suffix in {".pem", ".key"}:
            return True
        directory = self.root
        for i in range(len(path.parts)):
            local = Path(*path.parts[i:]).as_posix()
            if self._spec(directory).match_file(local):
                return True
            directory /= path.parts[i]
        return False

    def path(self, relative: str) -> Path:
        if self.ignored(relative):
            raise ValueError("Path is outside the repository or excluded from indexing")
        candidate = self.root / relative
        # Reject symlink components, including symlinks to files inside the repo.
        current = self.root
        for part in Path(relative).parts:
            current /= part
            if current.is_symlink():
                raise ValueError("Symlink paths are not indexed")
        resolved = candidate.resolve()
        if not resolved.is_relative_to(self.root) or resolved == self.root:
            raise ValueError("Path must identify a file inside the repository")
        return resolved

    def read(self, relative: str) -> str:
        path = self.path(relative)
        if not self.supported(relative):
            raise ValueError("File type is not authorized for this source")
        if path.stat().st_size > self.size_limit(relative):
            raise ValueError("File exceeds the indexing size limit")
        suffix = path.suffix.lower()
        if suffix in IMAGE_EXTENSIONS:
            with self.image(relative) as image:
                return f"Image: {relative}\nDimensions: {image.width} × {image.height} (preview)"
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted or len(reader.pages) > 200:
                raise ValueError("Encrypted PDFs and PDFs over 200 pages are unsupported")
            parts, count = [], 0
            for number, page in enumerate(reader.pages, 1):
                part = f"[Page {number}]\n{page.extract_text() or ''}"
                count += len(part)
                if count > MAX_BYTES:
                    raise ValueError("Extracted document exceeds text limit")
                parts.append(part)
            text = "\n\n".join(parts)
            if not any((part.split("\n", 1)[-1]).strip() for part in parts):
                raise ValueError("PDF contains no extractable text; OCR is not included")
        elif suffix == ".docx":
            # Bound uncompressed XML before letting the document parser expand it.
            import zipfile
            with zipfile.ZipFile(path) as archive:
                if sum(info.file_size for info in archive.infolist()) > MEDIA_MAX_BYTES:
                    raise ValueError("Expanded document exceeds size limit")
            from docx import Document
            document = Document(path)
            text = "\n".join([p.text for p in document.paragraphs] +
                             [" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows])
        else:
            text = path.read_text(encoding="utf-8")
        if len(text) > MAX_BYTES:
            raise ValueError("Extracted document exceeds text limit")
        if "\x00" in text:
            raise ValueError("Binary file")
        return text

    def fingerprint(self, relative: str, text: str | None = None) -> str:
        if self.kind(relative) == "images":
            path = self.path(relative)
            if not self.supported(relative) or path.stat().st_size > MEDIA_MAX_BYTES:
                raise ValueError("Image is not authorized or exceeds size limit")
            return hashlib.sha256(path.read_bytes()).hexdigest()
        return digest(self.read(relative) if text is None else text)

    def image(self, relative: str):
        from PIL import Image, ImageOps
        path = self.path(relative)
        if not self.supported(relative) or self.kind(relative) != "images" or path.stat().st_size > MEDIA_MAX_BYTES:
            raise ValueError("Image is not authorized or exceeds size limit")
        with Image.open(path) as original:
            if original.width * original.height > 40_000_000:
                raise ValueError("Image exceeds 40 megapixels")
            image = ImageOps.exif_transpose(original).convert("RGB")
            image.thumbnail((1024, 1024))
            return image

    def files(self):
        # Git handles nested .gitignore rules, ignored parents, and tracked files.
        result = subprocess.run(
            ["git", "-C", str(self.root), "ls-files", "--cached", "--others",
             "--exclude-standard", "-z"], capture_output=True, check=False,
        )
        if result.returncode == 0:
            candidates = sorted(set(result.stdout.decode().split("\x00")) - {""})
        else:
            candidates = []
            for directory, dirs, files in os.walk(self.root, followlinks=False):
                dirs[:] = [d for d in dirs if d not in EXCLUDED
                           and not (Path(directory) / d).is_symlink()
                           and not self.ignored((Path(directory) / d).relative_to(self.root).as_posix() + "/")]
                candidates.extend((Path(directory) / f).relative_to(self.root).as_posix()
                                  for f in sorted(files))
        for relative in candidates:
            path = Path(relative)
            if not self.supported(relative):
                continue
            try:
                candidate = self.path(relative)
                if candidate.is_file() and candidate.stat().st_size <= self.size_limit(relative):
                    yield relative
            except (OSError, ValueError):
                continue
